import io
import re
import zipfile
from pathlib import Path
from .config import MAX_CHUNKS


class ParseError(Exception):
    pass


class NeedsOCR(ParseError):
    pass


def parse(path: Path, extension: str):
    units = []
    if extension in {'docx', 'xlsx', 'pptx'}:
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 150 * 1024 * 1024 or len(z.infolist()) > 10000:
                raise ParseError('Office 文件解压后超过安全处理上限。')
    if extension == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ParseError('加密 PDF 无法解析，请上传解密后的副本。')
        if len(reader.pages) > 1000:
            raise ParseError('PDF 超过 1000 页处理上限。')
        empty = []
        for i, p in enumerate(reader.pages, 1):
            t = p.extract_text() or ''
            if len(t.strip()) < 10:
                empty.append(i)
            else:
                units.append((t, {'type': 'page', 'page': i, 'label': f'第 {i} 页'}))
        if not units:
            raise NeedsOCR('没有可提取文字，可能是扫描版 PDF；需要先进行 OCR。')
        warning = f'第 {", ".join(map(str,empty[:20]))} 页没有可提取文字，可能需要 OCR。' if empty else ''
    elif extension == 'docx':
        from docx import Document
        doc = Document(path)
        for i, p in enumerate(doc.paragraphs, 1):
            if p.text.strip():
                units.append((p.text, {'type': 'paragraph', 'paragraph': i, 'label': f'第 {i} 段'}))
        for ti, table in enumerate(doc.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                units.append((' | '.join(c.text for c in row.cells),
                              {'type': 'table', 'table': ti, 'row': ri, 'label': f'表 {ti} · 行 {ri}'}))
    elif extension == 'xlsx':
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        # Stored filenames are opaque IDs, so hand openpyxl a stream rather than a suffix-less path.
        wb = load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
        for ws in wb:
            if ws.max_row and ws.max_row > 50000 or ws.max_column and ws.max_column > 200:
                raise ParseError('工作表超过 50000 行或 200 列处理上限。')
            header = ''
            for ri, row in enumerate(ws.iter_rows(), 1):
                values = [f'{c.coordinate}={c.value}' for c in row if c.value is not None]
                if values:
                    text = ' | '.join(values)
                    if not header:
                        header = text
                    last = max(c.column for c in row if c.value is not None)
                    cell_range = f'A{ri}:{get_column_letter(last)}{ri}'
                    units.append((f'工作表 {ws.title}\n表头：{header}\n{text}',
                                  {'type': 'cells', 'sheet': ws.title, 'range': cell_range,
                                   'header_range': f'A1:{get_column_letter(last)}1',
                                   'label': f'{ws.title} · {cell_range}（表头第 1 行）'}))
        wb.close()
    elif extension == 'pptx':
        from pptx import Presentation
        pres = Presentation(path)
        for i, slide in enumerate(pres.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
            units.append(('\n'.join(texts), {'type': 'slide', 'slide': i, 'label': f'第 {i} 张幻灯片'}))
    else:
        try:
            text = path.read_text(encoding='utf-8-sig')
        except UnicodeDecodeError:
            text = path.read_text(encoding='gb18030')
        for i, paragraph in enumerate(re.split(r'\n\s*\n', text), 1):
            units.append((paragraph, {'type': 'paragraph', 'paragraph': i, 'label': f'第 {i} 段'}))
    chunks = []
    for text, locator in units:
        text = re.sub(r'[ \t]+', ' ', text).strip()
        # Overlap remains within its source unit, so locators never cross pages or rows.
        for start in range(0, len(text), 420):
            piece = text[start:start+500].strip()
            if piece:
                chunks.append((piece, {**locator, 'offset': start}))
            if len(chunks) > MAX_CHUNKS:
                raise ParseError(f'文档超过 {MAX_CHUNKS} 个片段上限，请拆分后上传。')
    if not chunks:
        raise ParseError('未找到可检索文字；空文档、图片和没有缓存值的公式不参与检索。')
    return chunks, locals().get('warning', '')
