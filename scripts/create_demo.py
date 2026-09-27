"""Build the fictional portfolio corpus without reading accounts or user files.

Generated files are committed so trying the project does not require regenerating
Office fixtures. This portable builder uses the project's locked Python libraries.
"""
import hashlib
import json
from datetime import datetime
from pathlib import Path
from demo_corpus import (COMPANY, FICTION, LIBRARIES, DOCUMENTS, TRAVEL_PAGES,
                         HR_SECTIONS, IT_PARAGRAPHS, PROJECT_SLIDES, EXPENSE_ROWS, WORKBOOKS)

ROOT=Path(__file__).resolve().parent.parent
DEMO=ROOT/'demo'
CREATED=datetime(2026,9,1)


def split_sections(text):
    blocks=text.strip().split('\n\n')
    result=[]
    for b in blocks:
        if b.startswith('# '): result.append(('title',b[2:]))
        elif b.startswith('## '): result.append(('heading',b[3:]))
        else: result.append(('body',b))
    return result


def write_docx(path, blocks):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml.ns import qn
    doc=Document()
    section=doc.sections[0]
    section.top_margin=section.bottom_margin=Cm(2)
    section.left_margin=section.right_margin=Cm(2.3)
    for name,pt in [('Normal',10.5),('Title',22),('Heading 1',15),('Heading 2',12)]:
        s=doc.styles[name]
        s.font.name='Microsoft YaHei'; s.font.size=Pt(pt); s.font.color.rgb=RGBColor(30,45,37)
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
        s.paragraph_format.space_after=Pt(8)
    doc.styles['Normal'].paragraph_format.line_spacing=1.35
    for kind,body in blocks:
        if kind=='title': doc.add_paragraph(body,'Title')
        elif kind=='heading': doc.add_heading(body,1)
        else: doc.add_paragraph(body)
    p=doc.add_paragraph('资料维护记录', 'Heading 1')
    table=doc.add_table(rows=1,cols=3); table.style='Light Shading Accent 1'
    for cell,value in zip(table.rows[0].cells,['日期','事项','维护角色']): cell.text=value
    for row in [('2026-09-01','发布演示版本','文档负责人'),('2026-09-22','补充流程 例外和案例','部门资料负责人')]:
        for cell,value in zip(table.add_row().cells,row): cell.text=value
    props=doc.core_properties
    props.author=props.last_modified_by='星桥科技虚构演示'
    props.created=props.modified=CREATED
    props.title=next((b for k,b in blocks if k=='title'),path.stem)
    props.subject='知序项目虚构语料'; props.comments=FICTION
    doc.save(path)


def write_pdf(path, pages=None, blocks=None):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.enums import TA_LEFT
    from xml.sax.saxutils import escape
    if 'STSong-Light' not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    title=ParagraphStyle('Title',fontName='STSong-Light',fontSize=20,leading=29,spaceAfter=15)
    heading=ParagraphStyle('Heading',fontName='STSong-Light',fontSize=14,leading=21,spaceBefore=9,spaceAfter=8)
    body=ParagraphStyle('Body',fontName='STSong-Light',fontSize=10.5,leading=17,spaceAfter=10,wordWrap='CJK',alignment=TA_LEFT)
    doc=SimpleDocTemplate(str(path),pagesize=(595,842),rightMargin=48,leftMargin=48,topMargin=56,bottomMargin=56,
                          title=path.stem,author='星桥科技虚构演示',subject=FICTION)
    story=[]
    if pages:
        for i,(name,paragraphs) in enumerate(pages):
            if i: story.append(PageBreak())
            story.append(Paragraph('出差管理办法' if i==0 else name,title))
            story.append(Paragraph('星桥科技 · FIN-TR-2026-01 · 2026.09 · 虚构演示',body))
            if i==0: story.append(Paragraph(FICTION,body))
            for p in paragraphs: story.append(Paragraph(escape(p),body))
    else:
        for kind,text in blocks:
            story.append(Paragraph(escape(text).replace('\n','<br/>'),title if kind=='title' else heading if kind=='heading' else body))
    def footer(canvas,document):
        canvas.setFont('STSong-Light',8)
        canvas.drawString(48,30,'星桥科技虚构演示资料 · 仅用于知序项目展示')
        canvas.drawRightString(547,30,f'第 {document.page} 页')
    doc.build(story,onFirstPage=footer,onLaterPages=footer)


def write_xlsx(path, sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb=Workbook(); wb.remove(wb.active)
    wb.properties.creator=wb.properties.lastModifiedBy='星桥科技虚构演示'
    wb.properties.created=wb.properties.modified=CREATED; wb.properties.description=FICTION
    for name,rows in sheets.items():
        ws=wb.create_sheet(name[:31])
        for row in rows: ws.append(row)
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:1'
        for row in ws:
            ws.row_dimensions[row[0].row].height=48 if name=='正文说明' else 42
            for cell in row:
                cell.font=Font(name='Microsoft YaHei',size=10,color='233E2B')
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if cell.row==1:
                    cell.fill=PatternFill('solid',fgColor='31573F'); cell.font=Font(name='Microsoft YaHei',size=10,color='FFFFFF',bold=True)
                elif cell.row%2==0: cell.fill=PatternFill('solid',fgColor='F0F5EA')
        for col in ws.columns:
            letter=col[0].column_letter
            ws.column_dimensions[letter].width=66 if name=='正文说明' and letter=='B' else 32 if letter!='A' else 25
        if name=='正文说明':
            for i in range(2,ws.max_row+1): ws.row_dimensions[i].height=110
    wb.save(path)


def write_pptx(path, slides):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    pres=Presentation(); pres.slide_width=Inches(13.333); pres.slide_height=Inches(7.5)
    props=pres.core_properties
    props.author=props.last_modified_by='星桥科技虚构演示'; props.created=props.modified=CREATED
    props.title=path.stem; props.subject=FICTION
    for i,(title,body) in enumerate(slides):
        slide=pres.slides.add_slide(pres.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor(246,249,241)
        box=slide.shapes.add_textbox(Inches(.8),Inches(.7),Inches(11.7),Inches(.8))
        p=box.text_frame.paragraphs[0]; p.text=title; p.font.name='Microsoft YaHei'; p.font.size=Pt(30); p.font.bold=True; p.font.color.rgb=RGBColor(36,72,48)
        box=slide.shapes.add_textbox(Inches(.85),Inches(1.85),Inches(11.6),Inches(4.6))
        tf=box.text_frame; tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.TOP
        for j,line in enumerate(body.split('\n')):
            p=tf.paragraphs[0] if j==0 else tf.add_paragraph(); p.text=line
            p.font.name='Microsoft YaHei'; p.font.size=Pt(22 if len(body)<330 else 18)
            p.font.color.rgb=RGBColor(57,74,61); p.space_after=Pt(20)
        box=slide.shapes.add_textbox(Inches(.85),Inches(6.83),Inches(11.6),Inches(.35))
        p=box.text_frame.paragraphs[0]; p.text=f'星桥科技虚构演示 · 2026.09     {i+1:02d} / {len(slides):02d}'
        p.font.name='Microsoft YaHei'; p.font.size=Pt(10); p.font.color.rgb=RGBColor(107,124,101)
    pres.save(path)


def create():
    for lib in LIBRARIES: (DEMO/lib['key']).mkdir(parents=True,exist_ok=True)
    entries=[]
    def add(path,library,directory,owner,code):
        entries.append({'path':path.relative_to(DEMO).as_posix(),'library':library,'directory':directory,'owner':owner,'code':code,
                        'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'size':path.stat().st_size})
    p=DEMO/'public/出差管理办法（虚构）.pdf'; write_pdf(p,pages=TRAVEL_PAGES); add(p,'public','行政与差旅','许澄','FIN-TR-2026-01')
    hr=[('title','星桥科技人事制度'),('body',FICTION),('heading','请假流程'),
        ('body','请假申请需提前3个工作日提交，填写请假类型、起止日期和交接人，由直属负责人审批。紧急情况应先联系负责人。'),
        ('heading','入职办理'),('body','新员工入职当天由人事演示专员安排账号开通和培训。办公设备由 IT 服务台登记发放。')]
    hr.extend(part for heading,text in HR_SECTIONS for part in [('heading',heading),('body',text)])
    p=DEMO/'public/人事制度（虚构）.docx'; write_docx(p,hr); add(p,'public','入职与人事','顾宁','HR-2026-09')
    p=DEMO/'public/费用标准（虚构）.xlsx'
    write_xlsx(p,{'费用标准':EXPENSE_ROWS,'审批矩阵':[('场景','发起角色','审批与复核','依据'),('常规出差','员工','直属负责人','填写申请与预算'),('超过5000元','员工','部门负责人及财务','核对预算与必要性'),('住宿超标','员工','负责人及财务书面确认','附报价与原因'),('培训超过2000元','员工','部门负责人及人事','岗位关联与成果'),('例外申请','员工','制度负责人确认','只覆盖本次申请')],
                       '案例记录':[('编号','角色','业务事项','材料','状态'),('TR-018','林晓','青禾访谈出差','申请、审批、发票、行程','材料完整'),('TR-021','叶遥','星舟培训出差','发票、支付截图','缺审批待补正'),('TR-022','沈言','客户承担住宿','交通和剩余餐食','住宿不重复申报'),('TR-023','林晓','会议改期','原票、新票、改期通知','核对实际损失')]}); add(p,'public','行政与差旅','许澄','FIN-EXP-2026')
    p=DEMO/'public/星舟项目介绍（虚构）.pptx'; write_pptx(p,PROJECT_SLIDES); add(p,'public','团队项目','林舟','PRJ-XZ-2026')
    p=DEMO/'public/办公设备指引（虚构）.txt'; p.write_text('\n\n'.join(IT_PARAGRAPHS),encoding='utf-8'); add(p,'public','办公与信息','陈屿','IT-USER-2026')
    access='# 虚构演示 · 星桥科技信息访问规范\n\n'+FICTION+'\n\n员工只能访问管理员授权的知识库；需要新增访问权限时，应联系知识库管理员说明用途。\n\n资料有冲突时，应向制度负责人确认适用版本，不能自行认定其中一份优先。\n\n'
    access+='## 资料分级与维护\n\n通用制度由对应部门维护，客户资料按项目角色使用，研发规划由研发负责人确认阅读范围。每份资料说明维护人、适用对象和版本。需求提出时包含使用目的和期限，负责人确认后调整。\n\n'
    access+='## 提交与复核\n\n员工提交的资料先交文库负责人核对来源、主题、版本和必要范围。尚未确认的草稿保留待审状态，批准后再进入正式目录。维护人每月检查过期内容、失效链接和同主题冲突，不把文件数量当作知识质量。\n\n'
    access+='## 客户和外部材料\n\n业务资料进入实际公司知识空间前按实际组织要求确认使用授权并清理敏感信息。公开演示全部使用虚构资料，不导入真实客户、真实员工身份材料或私人记录。原件位置与摘要保持可追溯关系。\n\n'
    access+='## 版本和撤回\n\n新版本记录修订人、日期和原因，旧版标记状态。撤回资料说明原因及替代文件；历史引用不能自动被视为现行制度。遇到同主题金额或日期冲突，向维护人确认而不是挑选看起来更方便的一份。\n\n'
    access+='## 示例与联系方式\n\n林晓申请阅读青禾交付资料，填写PRJ-QH-2026用途和预计期限；项目负责人确认范围后由管理员调整。演示咨询knowledge-owner@example.com。本文编号INFO-2026-09，维护人周砚，所有人物与联系信息均为虚构。'
    p=DEMO/'public/信息访问规范（虚构）.md'; p.write_text(access,encoding='utf-8'); add(p,'public','办公与信息','周砚','INFO-2026-09')
    private='# 虚构演示 · 研发保密空间\n\n'+FICTION+'\n\n演示保密代号：ORBIT-7429。研发演示发布日为2027年2月8日。仅研发保密库成员可访问。\n\n'
    private+='## 计划概况\n\nORBIT探索跨项目经验与来源链路，负责人林舟。先使用虚构样本核对相似案例的适用条件、版本和访问范围，再决定是否扩大资料覆盖。演示日期仅用于场景测试，不代表真实未公开项目。\n\n'
    private+='## 研发资料目录\n\n研发架构与技术决策记录系统分工及试点限制；ORBIT研发计划记录角色与节点；发布和回滚检查说明版本与恢复；风险与问题清单追踪资料冲突、扫描件和服务波动。各记录使用统一编号关联，不含生产命令或凭据。\n\n'
    private+='## 维护与交接\n\n林舟维护规划，宋禾维护质量与风险，陈屿维护运行与恢复，周砚维护业务假设。周会检查待确定事项和下一步动作，发布前核对记录与实际状态。未完成的研究不写成已交付能力。\n\n'
    private+='## 演示边界\n\n本文件为公开项目的虚构语料，名称中的保密仅用于应用授权场景，不包含真实保密资料。公开静态演示的虚构文件均可由访问者下载，不作为真实访问控制系统。完整服务端版本在查询、下载和引用处校验当前成员授权。'
    p=DEMO/'private/研发机密（虚构）.md'; p.write_text(private,encoding='utf-8'); add(p,'private','研发规划','林舟','RD-INDEX-001')
    for item in DOCUMENTS:
        p=DEMO/item['path']; blocks=split_sections(item['text'])
        if p.suffix=='.docx': write_docx(p,blocks)
        elif p.suffix=='.pdf': write_pdf(p,blocks=blocks)
        elif p.suffix=='.xlsx':
            sheets=dict(WORKBOOKS[item['library']]); sheets['正文说明']=[('主题','内容')]+[(kind,text) for kind,text in blocks]; write_xlsx(p,sheets)
        elif p.suffix=='.pptx':
            sections=[]; title='季度经营计划'; body=[]
            for kind,text in blocks:
                if kind in ('title','heading'):
                    if body: sections.append((title,'\n'.join(body))); body=[]
                    title=text
                else: body.append(text)
            if body: sections.append((title,'\n'.join(body)))
            write_pptx(p,sections)
        else: p.write_text(item['text'],encoding='utf-8')
        add(p,item['library'],item['directory'],item['owner'],item['code'])
    manifest={'version':2,'fictional':True,'company':COMPANY,'notice':FICTION,'libraries':LIBRARIES,'documents':entries}
    (DEMO/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'DEMO_READY: {len(entries)} documents in {len(LIBRARIES)} libraries; all fictional')


if __name__=='__main__': create()
